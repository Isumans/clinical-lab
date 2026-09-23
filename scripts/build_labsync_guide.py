from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "LabSync-DevOps-Learning-Guide.md"
OUTPUT = ROOT / "docs" / "LabSync-DevOps-Learning-Guide.docx"
LOGO = ROOT / "frontend" / "public" / "labsync-logo.png"

NAVY = "10243E"
BLUE = "31B4E7"
DARK_BLUE = "1F4D78"
TEXT = "24384A"
MUTED = "667D91"
LIGHT_BLUE = "E8F5FB"
LIGHT_GRAY = "F3F6F9"
BORDER = "CDDCE8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_first_row_as_header(table) -> None:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, text, fld_end])


def set_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("LabSync  |  Application, DevOps, and DevSecOps Guide")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_page_number(section.footer.paragraphs[0])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_header_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, default_color=TEXT) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, color=default_color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
            run.font.highlight_color = None
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, color=default_color, bold=True)
        else:
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, color=DARK_BLUE)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, color=default_color)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    mark_first_row_as_header(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EEF3F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color=NAVY)
        if index < len(lines) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    data = [rows[0]] + rows[2:]
    col_count = len(data[0])
    table = doc.add_table(rows=len(data), cols=col_count)
    widths = [9360 // col_count] * col_count
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    mark_first_row_as_header(table)
    table.style = "Table Grid"
    for r_index, row in enumerate(data):
        for c_index in range(col_count):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[c_index] if c_index < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, default_color=NAVY if r_index == 0 else TEXT)
            if r_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
            elif r_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    mark_first_row_as_header(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("NOTE  ")
    set_run_font(r, size=9.5, color=BLUE, bold=True)
    add_inline(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(
    doc: Document,
    title: str = "Application, DevOps, and\nDevSecOps Learning Guide",
    subtitle: str = "Combined technical project report, operational handbook,\nand beginner-friendly learning guide",
    tagline: str = "A documented journey from a minimum viable laboratory system to a secure, observable, recoverable cloud workload.",
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        picture = p.add_run().add_picture(str(LOGO), width=Inches(1.05))
        doc_pr = picture._inline.docPr
        doc_pr.set("name", "LabSync logo")
        doc_pr.set("descr", "LabSync laboratory management system logo")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("LABSYNC")
    set_run_font(r, size=12, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    set_run_font(r, size=28, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run(subtitle)
    set_run_font(r, size=13, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2700, 6660])
    mark_first_row_as_header(table)
    table.style = "Table Grid"
    metadata = [
        ("Release", "v0.3.4"),
        ("Application stack", "Next.js, TypeScript, FastAPI, SQLAlchemy, PostgreSQL"),
        ("Delivery platform", "GitHub Actions, GHCR, Docker Compose, Google Cloud"),
        ("Prepared", "August 2026"),
    ]
    for index, (label, value) in enumerate(metadata):
        set_cell_shading(table.cell(index, 0), LIGHT_BLUE)
        p1 = table.cell(index, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        set_run_font(r1, size=10, color=NAVY, bold=True)
        p2 = table.cell(index, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(value)
        set_run_font(r2, size=10, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(tagline)
    set_run_font(r, size=10.5, color=DARK_BLUE, italic=True)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_contents(doc: Document, headings: list[str]) -> None:
    p = doc.add_paragraph("Contents", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    for heading in headings:
        p = doc.add_paragraph(style="List Number")
        add_inline(p, re.sub(r"^\d+\.\s*", "", heading))
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def parse_markdown_body(doc: Document, lines: list[str]) -> None:
    index = 0
    paragraph_buffer: list[str] = []
    first_h1 = True

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer)
            p = doc.add_paragraph()
            add_inline(p, text)
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("```"):
            flush_paragraph()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index].rstrip("\n"))
                index += 1
            add_code_block(doc, code)
        elif line.startswith("| "):
            flush_paragraph()
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                table_rows.append(row)
                index += 1
            add_markdown_table(doc, table_rows)
            continue
        elif line.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
        elif line.startswith("## "):
            flush_paragraph()
            title = line[3:].strip()
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            doc.add_paragraph(title, style="Heading 1")
        elif line.startswith("# "):
            flush_paragraph()
        elif line.startswith("> "):
            flush_paragraph()
            add_callout(doc, line[2:].strip())
        elif re.match(r"^\d+\.\s+", line):
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif line.startswith("**") and line.endswith("  "):
            flush_paragraph()
        elif line.strip() in {"", "---"}:
            flush_paragraph()
        else:
            paragraph_buffer.append(line)
        index += 1
    flush_paragraph()


def build(
    source: Path = SOURCE,
    output: Path = OUTPUT,
    document_title: str = "LabSync: Application, DevOps, and DevSecOps Learning Guide",
    cover_title: str = "Application, DevOps, and\nDevSecOps Learning Guide",
    cover_subtitle: str = "Combined technical project report, operational handbook,\nand beginner-friendly learning guide",
    cover_tagline: str = "A documented journey from a minimum viable laboratory system to a secure, observable, recoverable cloud workload.",
) -> None:
    source_lines = source.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in source_lines if line.startswith("## ") and not line.startswith("## Appendix")]
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = document_title
    doc.core_properties.subject = "Combined technical project report and learning handbook"
    doc.core_properties.author = "LabSync Project"
    doc.core_properties.keywords = "LabSync, FastAPI, Next.js, PostgreSQL, Docker, CI/CD, GCP, DevSecOps"
    add_cover(doc, cover_title, cover_subtitle, cover_tagline)
    add_contents(doc, headings)

    body_start = next(i for i, line in enumerate(source_lines) if line.startswith("## 1."))
    parse_markdown_body(doc, source_lines[body_start:])

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
